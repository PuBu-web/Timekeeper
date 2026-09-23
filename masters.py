from __future__ import annotations

import difflib
import re
import unicodedata
import threading
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

FACTION_EMOJI = {"魏": "🔵", "蜀": "🟢", "呉": "🔴"}


def normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value or "")
    value = value.replace(" ", "").replace("　", "")
    return value.strip()


def row_value(row, index: int):
    """末尾が空欄の行でも安全に値を取得する。"""
    return row[index] if 0 <= index < len(row) else None


@dataclass(frozen=True)
class City:
    name: str
    reading: str
    x: int
    y: int
    category: str
    county: str
    faction: str

    @property
    def coordinate(self) -> str:
        return f"{self.x},{self.y}"

    @property
    def emoji(self) -> str:
        return FACTION_EMOJI.get(self.faction, "⚪")


@dataclass(frozen=True)
class GuildResolution:
    full_name: str
    short_name: str
    registered: bool
    faction: str = ""
    auto_added: bool = False
    save_error: str = ""
    learned: bool = False
    learning_error: str = ""


class ExcelMaster:
    """Version4.0の設定.xlsxを読み込み、保存時に自動リロードする。"""

    def __init__(self, xlsx_path: Path):
        self.xlsx_path = xlsx_path
        self.by_name: dict[str, City] = {}
        self.city_aliases: dict[str, str] = {}
        self.guilds: dict[str, tuple[str, str]] = {}
        self.guild_aliases: dict[str, str] = {}
        self._mtime = -1.0
        self._write_lock = threading.RLock()
        self.last_city_learning: tuple[str, str] | None = None
        self.last_city_learning_error = ""
        self.last_city_rejection_reason = ""
        self.reload(force=True)

    def reload(self, force: bool = False) -> bool:
        if not self.xlsx_path.exists():
            raise FileNotFoundError(f"Excelマスターが見つかりません: {self.xlsx_path}")
        mtime = self.xlsx_path.stat().st_mtime
        if not force and mtime == self._mtime:
            return False

        wb = load_workbook(self.xlsx_path, read_only=True, data_only=True)
        try:
            for required_sheet in ("城データ", "軍団データ"):
                if required_sheet not in wb.sheetnames:
                    raise RuntimeError(f"設定.xlsx に『{required_sheet}』シートがありません")

            city_ws = wb["城データ"]
            headers = {
                normalize_text(str(c.value or "")): i
                for i, c in enumerate(next(city_ws.iter_rows(min_row=1, max_row=1)))
            }
            required = ["城名", "X", "Y", "郡城", "国"]
            missing = [h for h in required if h not in headers]
            if missing:
                raise RuntimeError("城データの列が不足しています: " + ", ".join(missing))

            cities: dict[str, City] = {}
            for row in city_ws.iter_rows(min_row=2, values_only=True):
                name = normalize_text(str(row_value(row, headers["城名"]) or ""))
                if not name:
                    continue
                try:
                    x = int(row_value(row, headers["X"]))
                    y = int(row_value(row, headers["Y"]))
                except (TypeError, ValueError):
                    continue
                cities[name] = City(
                    name=name,
                    reading=normalize_text(str(row_value(row, headers.get("ふりがな", -1)) or "")) if "ふりがな" in headers else "",
                    x=x,
                    y=y,
                    category=normalize_text(str(row_value(row, headers.get("種類", -1)) or "")) if "種類" in headers else "",
                    county=normalize_text(str(row_value(row, headers["郡城"]) or "")),
                    faction=normalize_text(str(row_value(row, headers["国"]) or "")),
                )

            city_aliases: dict[str, str] = {}
            if "城OCR補正" in wb.sheetnames:
                ws = wb["城OCR補正"]
                h = {normalize_text(str(c.value or "")): i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
                if "OCR誤読城名" in h and "正式城名" in h:
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        wrong = normalize_text(str(row_value(row, h["OCR誤読城名"]) or ""))
                        correct = normalize_text(str(row_value(row, h["正式城名"]) or ""))
                        if wrong and correct:
                            city_aliases[wrong] = correct

            guild_ws = wb["軍団データ"]
            gheaders = {
                normalize_text(str(c.value or "")): i
                for i, c in enumerate(next(guild_ws.iter_rows(min_row=1, max_row=1)))
            }
            required_guild = ["正式軍団名", "登録名", "所属国"]
            missing_guild = [h for h in required_guild if h not in gheaders]
            if missing_guild:
                raise RuntimeError("軍団データの列が不足しています: " + ", ".join(missing_guild))

            guilds: dict[str, tuple[str, str]] = {}
            for row in guild_ws.iter_rows(min_row=2, values_only=True):
                full = self._core(str(row_value(row, gheaders["正式軍団名"]) or ""))
                short = normalize_text(str(row_value(row, gheaders["登録名"]) or ""))
                faction = normalize_text(str(row_value(row, gheaders["所属国"]) or ""))
                if full and short:
                    guilds[full] = (short, faction)

            guild_aliases: dict[str, str] = {}
            if "軍団OCR補正" in wb.sheetnames:
                ws = wb["軍団OCR補正"]
                h = {normalize_text(str(c.value or "")): i for i, c in enumerate(next(ws.iter_rows(min_row=1, max_row=1)))}
                if "OCR誤読名" in h and "正式軍団名" in h:
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        wrong = self._core(str(row_value(row, h["OCR誤読名"]) or ""))
                        correct = self._core(str(row_value(row, h["正式軍団名"]) or ""))
                        if wrong and correct:
                            guild_aliases[wrong] = correct

            self.by_name = cities
            self.city_aliases = city_aliases
            self.guilds = guilds
            self.guild_aliases = guild_aliases
            self._mtime = mtime
            return True
        finally:
            wb.close()

    @property
    def city_count(self) -> int:
        self.reload()
        return len(self.by_name)

    @property
    def guild_count(self) -> int:
        self.reload()
        return len(self.guilds)

    def find_city(self, raw_name: str, raw_coordinate: str = "") -> City | None:
        """Version4.1: 城名と座標を安全に照合し、見切れ推測の誤送信を防ぐ。"""
        self.reload()
        self.last_city_learning = None
        self.last_city_learning_error = ""
        self.last_city_rejection_reason = ""

        original_name = normalize_text(raw_name)
        name = re.sub(r"^(?:県城|郡城|港|関所|城)?Lv\d+", "", original_name, flags=re.IGNORECASE)
        name = re.sub(r"^(?:県城|郡城)", "", name)
        alias_corrected = self.city_aliases.get(name, name)

        coord = normalize_text(raw_coordinate).replace(".", ",")
        cm = re.fullmatch(r"(\d{1,4})\s*[,，]\s*(\d{1,4})", coord)
        coordinate_was_supplied = bool(coord)
        matched: City | None = None

        if cm:
            x, y = int(cm.group(1)), int(cm.group(2))
            exact = [c for c in self.by_name.values() if c.x == x and c.y == y]
            if len(exact) == 1:
                matched = exact[0]
            else:
                nearby = [c for c in self.by_name.values() if abs(c.x - x) <= 3 and abs(c.y - y) <= 3]
                if len(nearby) == 1:
                    matched = nearby[0]
                elif len(nearby) > 1:
                    ranked = sorted(
                        nearby,
                        key=lambda c: (
                            abs(c.x - x) + abs(c.y - y),
                            -difflib.SequenceMatcher(None, alias_corrected, c.name).ratio(),
                        ),
                    )
                    best, second = ranked[0], ranked[1]
                    best_key = (
                        abs(best.x - x) + abs(best.y - y),
                        -difflib.SequenceMatcher(None, alias_corrected, best.name).ratio(),
                    )
                    second_key = (
                        abs(second.x - x) + abs(second.y - y),
                        -difflib.SequenceMatcher(None, alias_corrected, second.name).ratio(),
                    )
                    if best_key < second_key:
                        matched = best

            # 座標が読めているのにマスター上の城が見つからない場合
            if matched is None:
                expected = self.by_name.get(alias_corrected)

                if expected is not None:
                    # Version4.1.2:
                    # 城名が正式名と完全一致していて、
                    # 座標の数字が1文字だけ違う場合はOCR誤読として正式座標へ補正する。
                    ocr_x = str(x)
                    ocr_y = str(y)
                    master_x = str(expected.x)
                    master_y = str(expected.y)

                    one_digit_coordinate_error = False

                    if len(ocr_x) == len(master_x) and len(ocr_y) == len(master_y):
                        diff_count = (
                            sum(a != b for a, b in zip(ocr_x, master_x))
                            + sum(a != b for a, b in zip(ocr_y, master_y))
                        )

                        if diff_count == 1:
                            one_digit_coordinate_error = True

                    if one_digit_coordinate_error:
                        matched = expected
                    else:
                        self.last_city_rejection_reason = (
                            f"城名『{name}』は登録済みですが、OCR座標 {coord} と "
                            f"正式座標 {expected.coordinate} が大きく異なります"
                        )
                        return None

                else:
                    self.last_city_rejection_reason = (
                        f"OCR座標 {coord} に一致・近接する城がありません"
                    )
                    return None
        elif coordinate_was_supplied:
            self.last_city_rejection_reason = f"座標形式が不完全です: {raw_coordinate}"
            return None
        else:
            # 座標が完全に空の場合のみ、登録済みの正式名・補正名を利用する。
            if alias_corrected in self.by_name:
                matched = self.by_name[alias_corrected]
            else:
                matches = difflib.get_close_matches(alias_corrected, self.by_name.keys(), n=1, cutoff=0.82)
                if matches:
                    matched = self.by_name[matches[0]]

        if matched is None:
            self.last_city_rejection_reason = f"城名『{name}』を安全に特定できません"
            return None

        # 座標から確定した城とOCR城名が著しく違う場合も、見切れ推測を疑って除外する。
        similarity = difflib.SequenceMatcher(None, alias_corrected, matched.name).ratio() if alias_corrected else 0.0
        if cm and alias_corrected and alias_corrected != matched.name and similarity < 0.45:
            self.last_city_rejection_reason = (
                f"座標は {matched.name}{matched.coordinate} ですが、OCR城名『{name}』との一致度が低すぎます"
            )
            return None

        # 学習は、座標で安全に確定できた場合、または座標なしで十分近い名前の場合だけ行う。
        if matched is not None and name and name != matched.name and name not in self.city_aliases:
            try:
                if self._append_learning("城OCR補正", "OCR誤読城名", "正式城名", name, matched.name):
                    self.last_city_learning = (name, matched.name)
            except Exception as exc:
                self.last_city_learning_error = str(exc)

        return matched

    @staticmethod
    def _core(name: str) -> str:
        name = normalize_text(name)
        name = re.sub(r"(?:軍軍団|軍団)$", "", name)
        name = re.sub(r"^.{1,6}?国", "", name)
        return name.strip()

    def resolve_guild(self, raw_name: str, auto_add: bool = True) -> GuildResolution:
        self.reload()
        core = self._core(raw_name)
        corrected = self.guild_aliases.get(core, core)

        if corrected in self.guilds:
            short, faction = self.guilds[corrected]
            return GuildResolution(corrected, short, True, faction=faction)

        registered = ""
        candidates = sorted(self.guilds, key=len, reverse=True)
        for candidate in candidates:
            if candidate in corrected or corrected in candidate:
                registered = candidate
                break

        if not registered:
            close = difflib.get_close_matches(corrected, self.guilds.keys(), n=1, cutoff=0.82)
            if close:
                registered = close[0]

        if registered:
            short, faction = self.guilds[registered]
            learned = False
            learning_error = ""
            if core and core != registered and core not in self.guild_aliases:
                try:
                    learned = self._append_learning(
                        "軍団OCR補正", "OCR誤読名", "正式軍団名", core, registered
                    )
                except Exception as exc:
                    learning_error = str(exc)
            return GuildResolution(
                registered, short, True, faction=faction,
                learned=learned, learning_error=learning_error,
            )

        fallback = corrected[:2] if corrected else "不明"
        if not auto_add or not corrected:
            return GuildResolution(corrected, fallback, False)

        try:
            self._append_unknown_guild(corrected, fallback)
            return GuildResolution(corrected, fallback, False, auto_added=True)
        except Exception as exc:
            return GuildResolution(corrected, fallback, False, save_error=str(exc))

    def _append_learning(
        self, sheet_name: str, wrong_header: str, correct_header: str, wrong: str, correct: str
    ) -> bool:
        """確定できたOCR誤読を補正シートへ重複なしで保存する。"""
        wrong = normalize_text(wrong)
        correct = normalize_text(correct)
        if not wrong or not correct or wrong == correct:
            return False

        with self._write_lock:
            wb = load_workbook(self.xlsx_path)
            try:
                if sheet_name not in wb.sheetnames:
                    ws = wb.create_sheet(sheet_name)
                    ws.append([wrong_header, correct_header, "メモ"])
                else:
                    ws = wb[sheet_name]

                headers = {normalize_text(str(c.value or "")): i + 1 for i, c in enumerate(ws[1])}
                if wrong_header not in headers or correct_header not in headers:
                    raise RuntimeError(f"{sheet_name}の列が不足しています")

                for row in ws.iter_rows(min_row=2, values_only=True):
                    existing_wrong = normalize_text(str(row[headers[wrong_header] - 1] or ""))
                    if existing_wrong == wrong:
                        return False

                values = [None] * max(len(headers), 3)
                values[headers[wrong_header] - 1] = wrong
                values[headers[correct_header] - 1] = correct
                if "メモ" in headers:
                    values[headers["メモ"] - 1] = "学習版が自動登録"
                ws.append(values)
                wb.save(self.xlsx_path)
            finally:
                wb.close()

        self.reload(force=True)
        return True

    def _append_unknown_guild(self, full_name: str, short_name: str) -> None:
        """未登録軍団を軍団データへ仮登録する。所属国は空欄で、後から確認する。"""
        wb = load_workbook(self.xlsx_path)
        try:
            ws = wb["軍団データ"]
            headers = {normalize_text(str(c.value or "")): i + 1 for i, c in enumerate(ws[1])}
            existing = {
                self._core(str(row[headers["正式軍団名"] - 1] or ""))
                for row in ws.iter_rows(min_row=2, values_only=True)
                if row
            }
            if full_name not in existing:
                ws.append([full_name, short_name, "", "未登録のため自動追加。所属国を入力してください"])
                wb.save(self.xlsx_path)
        finally:
            wb.close()
        self.reload(force=True)

    def short_guild_name(self, raw_name: str) -> str:
        return self.resolve_guild(raw_name, auto_add=False).short_name

    def guild_emoji(self, raw_name: str, fallback_faction: str = "") -> str:
        result = self.resolve_guild(raw_name, auto_add=False)
        return FACTION_EMOJI.get(result.faction or fallback_faction, "⚪")


class CityMaster:
    def __init__(self, xlsx_path: Path):
        self.master = ExcelMaster(xlsx_path)

    def find(self, raw_name: str, raw_coordinate: str = "") -> City | None:
        return self.master.find_city(raw_name, raw_coordinate)

    @property
    def last_learning(self) -> tuple[str, str] | None:
        return self.master.last_city_learning

    @property
    def last_learning_error(self) -> str:
        return self.master.last_city_learning_error

    @property
    def last_rejection_reason(self) -> str:
        return self.master.last_city_rejection_reason

    @property
    def count(self) -> int:
        return self.master.city_count


class GuildMaster:
    def __init__(self, xlsx_path: Path):
        self.master = ExcelMaster(xlsx_path)

    def resolve(self, raw_name: str, auto_add: bool = True) -> GuildResolution:
        return self.master.resolve_guild(raw_name, auto_add)

    def short_name(self, raw_name: str) -> str:
        return self.master.short_guild_name(raw_name)

    def emoji(self, raw_name: str, fallback_faction: str = "") -> str:
        return self.master.guild_emoji(raw_name, fallback_faction)

    @property
    def count(self) -> int:
        return self.master.guild_count
