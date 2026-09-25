from __future__ import annotations

import os
from pathlib import Path
from dotenv import load_dotenv
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")


def resolve_path(name: str, default: str) -> Path:
    raw = os.getenv(name, default).strip()
    path = Path(raw)
    return path if path.is_absolute() else BASE_DIR / path


def _env_int(name: str, default: int) -> int:
    raw = os.getenv(name, "").strip()
    if not raw:
        return default
    try:
        return int(raw)
    except ValueError:
        return default


def _to_int(value, default: int) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default

DISCORD_TOKEN = os.getenv("DISCORD_TOKEN", "").strip()
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
MASTER_XLSX_PATH = resolve_path("MASTER_XLSX_PATH", "data/設定.xlsx")
DB_PATH = resolve_path("DB_PATH", "data/siege_timekeeper_v4.db")
LOG_PATH = resolve_path("LOG_PATH", "logs/siege_timekeeper_v4.log")

# 初期値（設定.xlsxから上書き）
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.6-flash").strip()
IMAGE_INPUT_CHANNEL_ID = _env_int("IMAGE_INPUT_CHANNEL_ID", 1479759079227527199)
TIMEKEEPER_CHANNEL_ID = _env_int("TIMEKEEPER_CHANNEL_ID", 1223112232775323663)
LOG_CHANNEL_ID = _env_int("LOG_CHANNEL_ID", 0)
COOLDOWN_MINUTES = _env_int("COOLDOWN_MINUTES", 30)
OCCUPATION_DEDUPE_HOURS = _env_int("OCCUPATION_DEDUPE_HOURS", 24)
IMAGE_HASH_RETENTION_DAYS = _env_int("IMAGE_HASH_RETENTION_DAYS", 30)
DEBUG_MODE = os.getenv("DEBUG_MODE", "").strip().lower() not in {"0", "false", "off", "no"}
_SETTINGS_MTIME = -1.0


def reload_excel_settings(force: bool = False) -> bool:
    """設定.xlsxが更新されていれば設定シートを読み直す。環境変数がある項目は.envを優先。"""
    global _SETTINGS_MTIME
    global GEMINI_MODEL, IMAGE_INPUT_CHANNEL_ID, TIMEKEEPER_CHANNEL_ID, LOG_CHANNEL_ID
    global COOLDOWN_MINUTES, OCCUPATION_DEDUPE_HOURS, IMAGE_HASH_RETENTION_DAYS, DEBUG_MODE

    if not MASTER_XLSX_PATH.exists():
        return False
    mtime = MASTER_XLSX_PATH.stat().st_mtime
    if not force and mtime == _SETTINGS_MTIME:
        return False

    wb = load_workbook(MASTER_XLSX_PATH, read_only=True, data_only=True)
    try:
        if "設定" not in wb.sheetnames:
            raise RuntimeError("設定.xlsx に『設定』シートがありません")
        ws = wb["設定"]
        values = {}
        for row in ws.iter_rows(min_row=3, values_only=True):
            key = str(row[0]).strip() if row and row[0] is not None else ""
            if key:
                values[key] = row[1] if len(row) > 1 else None
    finally:
        wb.close()

    if not os.getenv("GEMINI_MODEL", "").strip():
        GEMINI_MODEL = str(values.get("GEMINI_MODEL") or GEMINI_MODEL).strip()
    if not os.getenv("IMAGE_INPUT_CHANNEL_ID", "").strip():
        IMAGE_INPUT_CHANNEL_ID = _to_int(values.get("IMAGE_INPUT_CHANNEL_ID"), IMAGE_INPUT_CHANNEL_ID)
    if not os.getenv("TIMEKEEPER_CHANNEL_ID", "").strip():
        TIMEKEEPER_CHANNEL_ID = _to_int(values.get("TIMEKEEPER_CHANNEL_ID"), TIMEKEEPER_CHANNEL_ID)
    if not os.getenv("LOG_CHANNEL_ID", "").strip():
        LOG_CHANNEL_ID = _to_int(values.get("LOG_CHANNEL_ID"), LOG_CHANNEL_ID)
    if not os.getenv("COOLDOWN_MINUTES", "").strip():
        COOLDOWN_MINUTES = _to_int(values.get("COOLDOWN_MINUTES"), COOLDOWN_MINUTES)
    if not os.getenv("OCCUPATION_DEDUPE_HOURS", "").strip():
        OCCUPATION_DEDUPE_HOURS = _to_int(values.get("OCCUPATION_DEDUPE_HOURS"), OCCUPATION_DEDUPE_HOURS)
    if not os.getenv("IMAGE_HASH_RETENTION_DAYS", "").strip():
        IMAGE_HASH_RETENTION_DAYS = _to_int(values.get("IMAGE_HASH_RETENTION_DAYS"), IMAGE_HASH_RETENTION_DAYS)
    if not os.getenv("DEBUG_MODE", "").strip():
        raw_debug = str(values.get("DEBUG_MODE") if values.get("DEBUG_MODE") is not None else "ON").strip().lower()
        DEBUG_MODE = raw_debug not in {"0", "false", "off", "no", "無効"}

    _SETTINGS_MTIME = mtime
    return True


reload_excel_settings(force=True)

ALLOWED_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
MAX_IMAGE_BYTES = 15 * 1024 * 1024
SPREADSHEET_ID = https://docs.google.com/spreadsheets/d/1r1ttDFCg24mEj6kFdvO_Lum19izous3Lul29Pq9v2y4/edit?gid=1830775209#gid=1830775209
